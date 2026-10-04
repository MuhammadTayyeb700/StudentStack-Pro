"""Slideshow engine: normalise AI output, build a .pptx, and build a standalone HTML presenter."""
from __future__ import annotations

import html
import io
import re

THEMES = {
    "Indigo": dict(bg="FFFFFF", hero="4F46E5", accent="7C3AED", text="1F2937", muted="6B7280", on_hero="FFFFFF"),
    "Ocean": dict(bg="FFFFFF", hero="0E7490", accent="0891B2", text="0F172A", muted="64748B", on_hero="FFFFFF"),
    "Forest": dict(bg="FFFFFF", hero="166534", accent="16A34A", text="14281D", muted="5B6B61", on_hero="FFFFFF"),
    "Sunset": dict(bg="FFFFFF", hero="C2410C", accent="F59E0B", text="27272A", muted="71717A", on_hero="FFFFFF"),
    "Midnight": dict(bg="0F172A", hero="1E293B", accent="38BDF8", text="E2E8F0", muted="94A3B8", on_hero="FFFFFF"),
}
LAYOUTS = ["bullets", "two_column", "section"]
MAX_SLIDES, MAX_BULLETS = 25, 7


def _s(v, limit: int) -> str:
    t = re.sub(r"\s+", " ", str(v or "")).strip()
    return t if len(t) <= limit else t[:limit].rsplit(" ", 1)[0].rstrip(",;:") + "…"


def _lines(v) -> list[str]:
    items = v if isinstance(v, list) else str(v or "").splitlines()
    out = [_s(re.sub(r"^\s*[-•*–]\s*", "", str(x)), 220) for x in items]
    return [x for x in out if x][:MAX_BULLETS]


def normalize(raw) -> dict:
    """Validate/repair a deck dict (from the AI or a restored backup). Raises ValueError if unusable."""
    if isinstance(raw, list):
        raw = {"slides": raw}
    if not isinstance(raw, dict):
        raise ValueError("The slideshow data is not in the expected format.")
    slides = []
    for s in (raw.get("slides") or [])[:MAX_SLIDES]:
        if not isinstance(s, dict):
            continue
        layout = s.get("layout") if s.get("layout") in LAYOUTS else "bullets"
        title, bullets, right = _s(s.get("title"), 90), _lines(s.get("bullets")), _lines(s.get("right"))
        if layout == "two_column" and not right:
            layout = "bullets"
        if not title and not bullets:
            continue
        slides.append(dict(layout=layout, title=title or "Untitled", bullets=bullets,
                           right=right if layout == "two_column" else [], notes=_s(s.get("notes"), 1200)))
    if not slides:
        raise ValueError("The AI didn't return any slides. Please try again.")
    return dict(title=_s(raw.get("title"), 90) or slides[0]["title"], subtitle=_s(raw.get("subtitle"), 140),
                slides=slides)


# ------------------------------------------------------------------ PowerPoint
def _size(bullets: list[str], cols: int = 1) -> int:
    n, total = len(bullets), sum(len(b) for b in bullets)
    lim = 1 if cols == 1 else 0.55
    if n <= 4 and total <= 260 * lim:
        return 26 if cols == 1 else 22
    if n <= 5 and total <= 380 * lim:
        return 22 if cols == 1 else 19
    if n <= 6 and total <= 520 * lim:
        return 20 if cols == 1 else 17
    return 18 if cols == 1 else 15


def to_pptx(deck: dict, theme: str = "Indigo", notes: bool = True) -> bytes:
    from lxml import etree
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
    from pptx.oxml.ns import qn
    from pptx.util import Inches, Pt

    deck, t = normalize(deck), THEMES.get(theme, THEMES["Indigo"])
    rgb = RGBColor.from_string
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    layout = prs.slide_layouts[5]  # "Title Only": gives every slide a real title placeholder (outline view, accessibility)

    def new_slide(bg_hex):
        sl = prs.slides.add_slide(layout)
        sl.background.fill.solid()
        sl.background.fill.fore_color.rgb = rgb(bg_hex)
        return sl

    def rect(sl, x, y, w, h, hex_):
        shp = sl.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
        shp.fill.solid()
        shp.fill.fore_color.rgb = rgb(hex_)
        shp.line.fill.background()
        sl.shapes._spTree.remove(shp._element)   # send behind the title placeholder
        sl.shapes._spTree.insert(2, shp._element)
        return shp

    def set_title(sl, text, x, y, w, h, size, color, bold=True, anchor=MSO_ANCHOR.MIDDLE):
        ph = sl.shapes.title
        ph.left, ph.top, ph.width, ph.height = x, y, w, h
        tf = ph.text_frame
        tf.word_wrap, tf.vertical_anchor = True, anchor
        tf.text = text
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.LEFT
        p.font.size, p.font.bold, p.font.color.rgb = Pt(size), bold, rgb(color)

    def textbox(sl, x, y, w, h, lines, size, color, bulleted=True, accent=None):
        tb = sl.shapes.add_textbox(x, y, w, h)
        tf = tb.text_frame
        tf.word_wrap = True
        for i, line in enumerate(lines):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.text = line
            p.font.size, p.font.color.rgb = Pt(size), rgb(color)
            p.space_after = Pt(size * 0.6)
            if bulleted:
                pPr = p._p.get_or_add_pPr()
                pPr.set("marL", "342900")
                pPr.set("indent", "-342900")
                bu = etree.SubElement(pPr, qn("a:buChar"))
                bu.set("char", "•")
        return tb

    def put_notes(sl, text):
        if notes and text:
            sl.notes_slide.notes_text_frame.text = text

    # title slide
    sl = new_slide(t["hero"])
    rect(sl, Inches(0.8), Inches(3.55), Inches(1.6), Inches(0.09), t["accent"])
    set_title(sl, deck["title"], Inches(0.8), Inches(1.5), Inches(11.7), Inches(2.0),
              44 if len(deck["title"]) <= 45 else 34, t["on_hero"], anchor=MSO_ANCHOR.BOTTOM)
    if deck["subtitle"]:
        textbox(sl, Inches(0.8), Inches(3.8), Inches(11.7), Inches(1.4), [deck["subtitle"]], 22, t["on_hero"], bulleted=False)

    for n, s in enumerate(deck["slides"], 2):
        if s["layout"] == "section":
            sl = new_slide(t["hero"])
            rect(sl, Inches(0.8), Inches(4.2), Inches(1.6), Inches(0.09), t["accent"])
            set_title(sl, s["title"], Inches(0.8), Inches(2.2), Inches(11.7), Inches(2.0), 38, t["on_hero"],
                      anchor=MSO_ANCHOR.BOTTOM)
            if s["bullets"]:
                textbox(sl, Inches(0.8), Inches(4.5), Inches(11.7), Inches(1.8), s["bullets"][:3], 20, t["on_hero"], bulleted=False)
        else:
            sl = new_slide(t["bg"])
            rect(sl, 0, 0, Inches(0.28), prs.slide_height, t["accent"])
            set_title(sl, s["title"], Inches(0.8), Inches(0.45), Inches(11.8), Inches(1.1),
                      32 if len(s["title"]) <= 50 else 26, t["text"])
            if s["layout"] == "two_column":
                sz = min(_size(s["bullets"], 2), _size(s["right"], 2))
                textbox(sl, Inches(0.8), Inches(1.9), Inches(5.8), Inches(4.8), s["bullets"], sz, t["text"])
                textbox(sl, Inches(6.9), Inches(1.9), Inches(5.8), Inches(4.8), s["right"], sz, t["text"])
            else:
                textbox(sl, Inches(0.8), Inches(1.9), Inches(11.7), Inches(4.9), s["bullets"], _size(s["bullets"]), t["text"])
            textbox(sl, Inches(12.2), Inches(6.9), Inches(0.8), Inches(0.4), [str(n)], 12, t["muted"], bulleted=False)
        put_notes(sl, s["notes"])
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


# ------------------------------------------------------------------ HTML presenter
def to_html(deck: dict, theme: str = "Indigo") -> str:
    """Self-contained slideshow: arrow keys / click / buttons, N toggles speaker notes, F toggles fullscreen."""
    deck, t = normalize(deck), THEMES.get(theme, THEMES["Indigo"])
    e = html.escape

    def ul(items):
        return "<ul>" + "".join(f"<li>{e(i)}</li>" for i in items) + "</ul>" if items else ""

    parts = [f'<section class="s hero"><h1>{e(deck["title"])}</h1><div class="bar"></div><p>{e(deck["subtitle"])}</p></section>']
    notes = [""]
    for s in deck["slides"]:
        if s["layout"] == "section":
            parts.append(f'<section class="s hero"><h2 class="big">{e(s["title"])}</h2><div class="bar"></div>{ul(s["bullets"][:3])}</section>')
        elif s["layout"] == "two_column":
            parts.append(f'<section class="s"><h2>{e(s["title"])}</h2><div class="cols"><div>{ul(s["bullets"])}</div>'
                         f'<div>{ul(s["right"])}</div></div></section>')
        else:
            parts.append(f'<section class="s"><h2>{e(s["title"])}</h2>{ul(s["bullets"])}</section>')
        notes.append(s["notes"])
    notes_html = "".join(f'<div class="n">{e(n) or "<i>No speaker notes</i>"}</div>' for n in notes)
    return f"""<!doctype html><html lang="en"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>{e(deck['title'])}</title><style>
*{{box-sizing:border-box}}html,body{{margin:0;height:100%;background:#000;overflow:hidden;font-family:Segoe UI,Helvetica,Arial,sans-serif}}
#wrap{{position:absolute;inset:0;display:flex;align-items:center;justify-content:center}}
#stage{{width:1280px;height:720px;position:relative;background:#{t['bg']};color:#{t['text']};transform-origin:center;flex:none;overflow:hidden}}
.s{{position:absolute;inset:0;padding:70px 90px;display:none}}.s.on{{display:block}}
.s::before{{content:"";position:absolute;left:0;top:0;bottom:0;width:16px;background:#{t['accent']}}}
.s.hero{{background:#{t['hero']};color:#{t['on_hero']};display:none;padding-top:200px}}.s.hero.on{{display:block}}.s.hero::before{{display:none}}
h1{{font-size:64px;margin:0 0 20px;line-height:1.1}}h2{{font-size:46px;margin:0 0 36px;line-height:1.15}}h2.big{{font-size:58px}}
.bar{{width:110px;height:7px;background:#{t['accent']};margin:0 0 24px}}.hero p{{font-size:30px;margin:0;opacity:.92}}
ul{{margin:0;padding-left:34px;font-size:32px;line-height:1.4}}li{{margin-bottom:14px}}.cols{{display:grid;grid-template-columns:1fr 1fr;gap:50px}}.cols ul{{font-size:27px}}
#ctl{{position:absolute;left:0;right:0;bottom:0;display:flex;gap:10px;align-items:center;justify-content:center;padding:8px;z-index:5;
background:linear-gradient(transparent,rgba(0,0,0,.55))}}
button{{background:rgba(255,255,255,.18);color:#fff;border:0;border-radius:8px;padding:8px 14px;font-size:15px;cursor:pointer}}button:hover{{background:rgba(255,255,255,.3)}}
#cnt{{color:#fff;font-size:14px;min-width:60px;text-align:center}}
#prog{{position:absolute;left:0;top:0;height:5px;background:#{t['accent']};width:0;z-index:6;transition:width .25s}}
#notes{{position:absolute;left:0;right:0;bottom:46px;background:rgba(15,23,42,.94);color:#e2e8f0;padding:14px 20px;font-size:16px;line-height:1.5;display:none;z-index:5;max-height:35%;overflow:auto}}
#notes.on{{display:block}}.n{{display:none}}.n.on{{display:block}}
</style></head><body><div id="wrap"><div id="stage">{''.join(parts)}<div id="prog"></div><div id="notes">{notes_html}</div>
<div id="ctl"><button id="p">◀</button><span id="cnt"></span><button id="x">▶</button><button id="nt">Notes (N)</button><button id="fs">Fullscreen (F)</button></div></div></div>
<script>
const S=[...document.querySelectorAll('.s')],N=[...document.querySelectorAll('.n')],st=document.getElementById('stage');let i=0;
function show(k){{i=Math.max(0,Math.min(S.length-1,k));S.forEach((e,j)=>e.classList.toggle('on',j==i));N.forEach((e,j)=>e.classList.toggle('on',j==i));
document.getElementById('cnt').textContent=(i+1)+' / '+S.length;document.getElementById('prog').style.width=(100*(i+1)/S.length)+'%';}}
function fit(){{const k=Math.min(innerWidth/1280,innerHeight/720);st.style.transform='scale('+k+')';}}
function fs(){{if(document.fullscreenElement)document.exitFullscreen();else document.documentElement.requestFullscreen&&document.documentElement.requestFullscreen().catch(()=>{{}});}}
document.getElementById('p').onclick=()=>show(i-1);document.getElementById('x').onclick=()=>show(i+1);
document.getElementById('nt').onclick=()=>document.getElementById('notes').classList.toggle('on');document.getElementById('fs').onclick=fs;
addEventListener('keydown',ev=>{{const k=ev.key;if(['ArrowRight','PageDown',' ','Enter'].includes(k)){{ev.preventDefault();show(i+1)}}
else if(['ArrowLeft','PageUp','Backspace'].includes(k))show(i-1);else if(k=='Home')show(0);else if(k=='End')show(S.length-1);
else if(k=='n'||k=='N')document.getElementById('notes').classList.toggle('on');else if(k=='f'||k=='F')fs();}});
addEventListener('resize',fit);fit();show(0);
</script></body></html>"""
